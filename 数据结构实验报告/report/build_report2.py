# -*- coding: utf-8 -*-
"""在已填好实验一的报告上追加实验二（模板 docx 已被移走，改为按文字定位锚点）。"""
import copy, os, shutil
from docx import Document
from docx.shared import Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn

BASE = r'E:\CODE\数据结构实验报告\report'
FIG = os.path.join(BASE, 'fig')
OUT = os.path.join(BASE, '数据结构与算法实验报告-郝永健-20252082264.docx')
BAK = os.path.join(BASE, '数据结构与算法实验报告-郝永健-20252082264.备份.docx')
XS = '{http://www.w3.org/XML/1998/namespace}space'

shutil.copyfile(OUT, BAK)
d = Document(OUT)
paras = list(d.paragraphs)


def find(text, start=0, exact=True):
    for i in range(start, len(paras)):
        t = paras[i].text.strip()
        if (t == text) if exact else t.startswith(text):
            return i
    raise KeyError(text)


i2 = find('实验二', exact=False)
i3 = find('实验三', exact=False)
i_intro_h = find('1内容简介', i2)
i_q_h = find('2 实验上机试题', i2)
i_h3 = find('3 算法实现', i2)
i_h4 = find('4 测试', i2)
i_h5 = find('5 难点和解决办法', i2)
i_tail = find('编写方法及要求请参考实验一', i2)
i_q1 = find('（1）', i2)

MODEL_BODY = paras[82]
MODEL_CAP = paras[77]
MODEL_H3 = next(p for p in paras if p.text.strip() == '3.1 算法流程设计')
MODEL_H32 = next(p for p in paras if p.text.strip() == '3.2 实现结果')
TBL_MODEL = copy.deepcopy(d.tables[0]._tbl)


def clone_para(model, text):
    new = copy.deepcopy(model._p)
    runs = new.findall(qn('w:r'))
    r_model = copy.deepcopy(runs[0]) if runs else None
    for ch in list(new):
        if ch.tag != qn('w:pPr'):
            new.remove(ch)
    if text:
        r = r_model if r_model is not None else new.makeelement(qn('w:r'), {})
        for ch in list(r):
            if ch.tag != qn('w:rPr'):
                r.remove(ch)
        rPr = r.find(qn('w:rPr'))
        if rPr is not None:
            for c in rPr.findall(qn('w:color')):
                rPr.remove(c)
        t = r.makeelement(qn('w:t'), {})
        t.text = text
        t.set(XS, 'preserve')
        r.append(t)
        new.append(r)
    return new


def body(text):
    return clone_para(MODEL_BODY, text)


def caption(text):
    return clone_para(MODEL_CAP, text)


def picture(name, width_cm):
    p = d.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.first_line_indent = Cm(0)
    p.paragraph_format.left_indent = Cm(0)
    p.add_run().add_picture(os.path.join(FIG, name + '.png'), width=Cm(width_cm))
    return p._p


def insert(anchor, elements):
    for el in elements:
        anchor._p.addprevious(el)


def set_cell(tc, text):
    ps = tc.findall(qn('w:p'))
    for extra in ps[1:]:
        tc.remove(extra)
    p = ps[0]
    runs = p.findall(qn('w:r'))
    for extra in runs[1:]:
        p.remove(extra)
    if runs:
        r = runs[0]
    else:
        r = p.makeelement(qn('w:r'), {})
        p.append(r)
    for ch in list(r):
        if ch.tag != qn('w:rPr'):
            r.remove(ch)
    t = r.makeelement(qn('w:t'), {})
    t.text = text
    t.set(XS, 'preserve')
    r.append(t)


def make_table(rows):
    tbl = copy.deepcopy(TBL_MODEL)
    trs = tbl.findall(qn('w:tr'))
    data_model = trs[1]
    for tr in trs[1:]:
        tbl.remove(tr)
    for tc, val in zip(trs[0].findall(qn('w:tc')), rows[0]):
        set_cell(tc, val)
    for row in rows[1:]:
        tr = copy.deepcopy(data_model)
        tbl.append(tr)
        for tc, val in zip(tr.findall(qn('w:tc')), row):
            set_cell(tc, val)
    return tbl


HEAD = ['用例编号', '测试项', '测试步骤', '预期结果', '测试结果']

# ---------- 1 内容简介 ----------
insert(paras[i_q_h], [body(
    '栈和队列是两种操作受限的线性表。栈只允许在表的一端（栈顶）进行插入和删除，特点是后进先出。顺序栈用一个数组和栈顶'
    '下标top实现，top初始为-1，入栈时先让top加1再存元素，出栈时先取出元素再让top减1，入栈、出栈和取栈顶的时间复杂度'
    '都是O(1)。栈的典型应用有表达式求值、括号匹配和函数调用。队列只允许在表的一端（队尾）插入、在另一端（队头）删除，'
    '特点是先进先出。为了充分利用数组空间，通常采用循环队列，让front指向队头元素、rear指向队尾元素的下一个位置，队空的'
    '条件是front==rear，队满的条件是(rear+1)%maxSize==front，这样需要浪费一个存储单元来区分队空和队满。队列常用于'
    '二项式系数（杨辉三角）的生成、广度优先搜索和缓冲区管理。串是由零个或多个字符组成的有限序列，可以采用定长顺序存储；'
    '模式匹配是串最重要的运算之一，朴素的BF算法让模式串从主串的每个位置开始逐字符比较，最坏时间复杂度为O(n×m)，KMP'
    '算法利用模式串自身的部分匹配信息，把时间复杂度降到O(n+m)。')])

# ---------- 2 实验上机试题 ----------
insert(paras[i_h3], [
    body('（1）设计一个可编程计算器，可以解析操作数为一位整数的四则运算表达式，如：3×(7-2×2)+8÷4。'),
    body('（2）利用队列辅助存储结构，计算(a+b)^k（k>=2）展开后各项系数（按照a的降幂、b的升幂顺序输出各项系数）。'),
    body('（3）设计一个算法，在串str中查找字符串substr最后一次出现的位置（不能使用标准模板库中的函数）。')])

# ---------- 3.1 / 3.2 ----------
H31 = clone_para(MODEL_H3, '3.1 算法流程设计')
H32 = clone_para(MODEL_H32, '3.2 实现结果')
s31 = [H31]
s31 += [body(
    '试题1要求解析操作数为一位整数的四则运算表达式。程序采用双栈法：一个操作数栈num存放数字和中间结果，一个运算符栈op'
    '存放运算符和左括号。从左到右扫描表达式，遇到数字直接压入num；遇到左括号压入op；遇到右括号则不断弹出op栈顶的运算符'
    '并计算，直到弹出左括号为止；遇到运算符时，只要op栈顶不是左括号且优先级不低于当前运算符，就先弹出它并计算，然后把'
    '当前运算符压入op。扫描结束后，把op中剩余的运算符依次弹出并计算，num栈中剩下的唯一元素就是结果。若出现括号不匹配或'
    '除数为0，程序给出错误提示。算法流程如图2.1所示。')]
s31 += [picture('fig2_1', 12.5), caption('图2.1 试题1算法流程图')]
s31 += [body(
    '试题2要求输出(a+b)^k展开后的各项系数。由二项式定理，各项系数就是杨辉三角第k行的k+1个数，并且满足'
    'C(k,i)=C(k-1,i-1)+C(k-1,i)。程序用循环队列保存当前一行的系数：初始时入队1、1（即k=1的那一行），然后从level=2'
    '开始逐层生成。生成新的一行时，先出队一个元素a，取新的队头b，把a+b入队，重复level-1次，最后再入队一个1，队列中'
    '就得到了第level行的全部系数。循环结束后依次出队输出即可。队列中始终只保存当前一行，辅助空间为O(k)。算法流程如'
    '图2.2所示。')]
s31 += [picture('fig2_2', 12.5), caption('图2.2 试题2算法流程图')]
s31 += [body(
    '试题3要求在串str中查找子串substr最后一次出现的位置，且不能使用标准模板库中的函数。程序先用自己编写的Length函数'
    '求出两个串的长度n和m，然后从主串中可能匹配的最后一个起点i=n-m开始向前逐个位置尝试：对每个i，用j从0到m-1逐字符'
    '比较str[i+j]与substr[j]，若m个字符全部相等，则i+1就是该子串出现的位置；由于是从后向前找，第一个找到的就是最后'
    '一次出现的位置。若所有起点都不匹配，则说明子串不存在。算法流程如图2.3所示。')]
s31 += [picture('fig2_3', 12.5), caption('图2.3 试题3算法流程图')]
s31 += [H32, body('三个程序在g++环境下编译运行，结果如图2.4至图2.12所示。')]
s31 += [body('试题1：输入含×、÷的表达式时程序能正确识别并计算，嵌套括号和除数为0的情况也能正确处理，如图2.4～'
              '图2.6所示。')]
for n, c in (('e2_1a', '图2.4 试题1实验结果1（含×、÷的表达式）'),
             ('e2_1b', '图2.5 试题1实验结果2（嵌套括号）'),
             ('e2_1c', '图2.6 试题1实验结果3（除数为0）')):
    s31 += [picture(n, 14.5), caption(c)]
s31 += [body('试题2：k=4和k=2时输出的系数与二项式定理一致，k小于2时给出提示，如图2.7～图2.9所示。')]
for n, c in (('e2_2a', '图2.7 试题2实验结果1（k=4）'),
             ('e2_2b', '图2.8 试题2实验结果2（k=2）'),
             ('e2_2c', '图2.9 试题2实验结果3（k=1，错误输入）')):
    s31 += [picture(n, 14.5), caption(c)]
s31 += [body('试题3：一般情况、子串重叠出现的情况以及子串不存在的情况都能正确处理，如图2.10～图2.12所示。')]
for n, c in (('e2_3a', '图2.10 试题3实验结果1（一般情况）'),
             ('e2_3b', '图2.11 试题3实验结果2（子串重叠）'),
             ('e2_3c', '图2.12 试题3实验结果3（子串不存在）')):
    s31 += [picture(n, 14.5), caption(c)]
insert(paras[i_h4], s31)

# ---------- 4 测试 ----------
T21 = [HEAD,
       ['1', '正常值', '输入3×(7-2×2)+8÷4', '输出11', '通过'],
       ['2', '正常值', '输入(1+2)*(3+4)', '输出21', '通过'],
       ['3', '边界值', '输入5（单个操作数）', '输出5', '通过'],
       ['4', '边界值', '输入8÷4', '输出2', '通过'],
       ['5', '错误值', '输入3*(7-2（括号不匹配）', '提示表达式错误', '通过'],
       ['6', '错误值', '输入3/0（除数为0）', '提示除数不能为0', '通过']]
T22 = [HEAD,
       ['1', '正常值', 'k=4', '输出1 4 6 4 1', '通过'],
       ['2', '正常值', 'k=6', '输出1 6 15 20 15 6 1', '通过'],
       ['3', '边界值', 'k=2', '输出1 2 1', '通过'],
       ['4', '边界值', 'k=3', '输出1 3 3 1', '通过'],
       ['5', '错误值', 'k=1（k小于2）', '提示k必须大于等于2', '通过']]
T23 = [HEAD,
       ['1', '正常值', 'str=ababcabcab，substr=ab', '输出9', '通过'],
       ['2', '边界值', 'str=aaaa，substr=aa（重叠出现）', '输出3', '通过'],
       ['3', '边界值', 'str=abc，substr=abc（整串匹配）', '输出1', '通过'],
       ['4', '边界值', 'str=abc，substr=abcd（子串更长）', '提示不存在', '通过'],
       ['5', '错误值', 'str=hello，substr=xyz', '提示不存在', '通过']]
s4 = [body('针对每个试题，分别按正常值、边界值和错误值设计测试用例，各用例的测试步骤与结果见表2.1至表2.3。')]
for i, T in enumerate((T21, T22, T23), start=1):
    s4 += [caption('表2.%d 试题%d测试用例' % (i, i)), make_table(T)]
insert(paras[i_h5], s4)

# ---------- 5 难点和解决办法 ----------
s5 = [body(
    '试题1的难点是运算符优先级的处理和括号的匹配。最初想把中缀表达式先转换成后缀表达式再计算，但这样要扫描两遍，代码也'
    '更长；后来改用两个栈直接边扫描边计算：遇到运算符时，只要栈顶运算符的优先级不低于当前运算符，就说明栈顶的运算可以'
    '先做，弹出并计算，直到栈顶是左括号或优先级更低为止，再把当前运算符入栈。右括号的处理是不断弹出运算符并计算，直到'
    '弹出左括号；如果栈里根本找不到左括号，就说明括号不匹配。扫描结束后还要把运算符栈清空，如果清空过程中遇到了左括号，'
    '同样说明括号不匹配。除数为0的情况单独判断，否则程序会得到无穷大而不报错。'),
      body(
    '另一个不容易发现的问题是×和÷的输入。题目中的例子用的是全角的乘号、除号，而Windows控制台默认使用GBK编码，源程序'
    '文件是UTF-8编码，同一个符号在两种编码下占用的字节并不相同：×在UTF-8中是C3 97，在GBK中是A1 C1；÷在UTF-8中是'
    'C3 B7，在GBK中是A1 C2。如果只按单个char比较，这两个符号会被当成两个非法字符，表达式就无法计算。程序中写了NextChar'
    '函数，同时识别这两种编码的双字节序列，把它们统一转换成ASCII的*和/，这样无论控制台用哪种编码输入，程序都能正确'
    '处理。'),
      body(
    '试题2的关键是理解队列中相邻两个元素之和正好是下一行的系数。初始队列中放的是第1行的1、1，生成第level行时，先把队头'
    '元素出队保存为a，此时新的队头b就是a右边相邻的元素，a+b正好是下一行对应位置的系数，把它入队；重复level-1次后再入队'
    '一个1，就补齐了新一行的首尾两个1。这里容易出错的是循环次数：第level行由上一行的level个元素生成，需要做level-1次'
    '相加，循环写成i<=level-1。另外循环队列的判满条件是(rear+1)%maxSize==front，会浪费一个存储单元，所以队列容量要'
    '开到k+2，否则生成最后一行时会溢出。'),
      body(
    '试题3要求不能使用标准模板库中的函数，因此字符串长度也要自己用循环统计。为了找到最后一次出现的位置，最直接的做法是'
    '从左到右扫描并记录每一次匹配的起点，但这样即使已经找到也要把整个主串扫完；改成从i=n-m开始从右向左扫描后，找到的'
    '第一个匹配就是最后一次出现，可以提前结束。需要特别注意子串重叠的情况，例如str="aaaa"、substr="aa"时，位置1、2、3'
    '都匹配，最后一次是3，从右向左扫描正好能自然处理。位置按从1开始编号，与题目中“位置”的习惯一致。'),
      body(
    '这三道题涉及的结构在实际中都很常见：表达式求值是编译器和计算器的核心功能，双栈法也是很多脚本语言解释器的实现基础；'
    '队列生成二项式系数的思路可以推广到任何需要“用上一行推下一行”的递推问题，只保留当前一行可以节省大量空间；串的模式'
    '匹配则用于文本编辑器的查找替换、日志分析和生物序列比对，理解BF算法之后再学习KMP就更容易体会“利用已匹配信息”的'
    '意义。')]
insert(paras[i_tail], s5)

# ---------- 删除占位 ----------
for p in (paras[i_q1], paras[i_q1 + 1], paras[i_q1 + 2],
          paras[find('此部分写近两周所学内容', i2)]):
    p._p.getparent().remove(p._p)
paras[i_tail]._p.getparent().remove(paras[i_tail]._p)

d.save(OUT)
print('saved', OUT)
